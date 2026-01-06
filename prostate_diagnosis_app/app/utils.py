import os
import pickle
import numpy as np
import tensorflow as tf
from io import BytesIO
from docx import Document
from docx.shared import Inches
from tensorflow.keras.models import load_model
from flask import current_app
import joblib

# 📍 Paths to the scaler and model files (relative to app root)
SCALER_PATH = os.path.join('app', 'ml', 'scaler_ann.pkl')
MODEL_PATH = os.path.join('app', 'ml', 'ann_model(new).keras')

# ✅ Load the scaler
try:
    scaler = joblib.load(SCALER_PATH)
except FileNotFoundError:
    raise RuntimeError(f"Scaler file not found at {SCALER_PATH}")
except Exception as e:
    raise RuntimeError(f"Error loading scaler: {e}")

# ✅ Load the ANN model
try:
    model = load_model(MODEL_PATH)
except FileNotFoundError:
    raise RuntimeError(f"Model file not found at {MODEL_PATH}")
except Exception as e:
    raise RuntimeError(f"Error loading model: {e}")


try:
    model = load_model(MODEL_PATH)
except OSError:
    raise RuntimeError(f"ANN model file not found at {MODEL_PATH}")


# Define a mapping if you're not using LabelEncoder
RACE_MAPPING = {
    'White': 0,
    'Black': 1,
    'Asian': 2,
    'Hispanic': 3,
    'Other': 4
}

def preprocess_and_predict(form_data):
    """
    Takes raw form data, preprocesses it using the scaler,
    and returns the predicted label and probability score.
    """
    try:
        features = [
            int(form_data['age']),
            RACE_MAPPING[form_data['race']],  # 🔄 string to int
            float(form_data['bmi']),
            float(form_data['psa_level']),
            float(form_data['prostate_volume']),
            int(form_data['family_history']),
            int(form_data['dre_result']),
            int(form_data['urinary_frequency']),
            int(form_data['nocturia']),
            int(form_data['weak_urine_stream']),
            int(form_data['hematuria']),
            int(form_data['erectile_dysfunction']),
            int(form_data['pain_in_pelvis']),
            int(form_data['bone_pain']),
            int(form_data['weight_loss'])
        ]
    except (KeyError, ValueError) as e:
        raise ValueError(f"Invalid form data: {e}")
    
    X = scaler.transform([features])
    prediction = model.predict(X)[0][0]
    label = "Positive" if prediction > 0.5 else "Negative"
    confidence = prediction if label == "Positive" else 1 - prediction
    return label, float(confidence)



def generate_docx_report(context):
    """
    Generates a Word document report (.docx) for the diagnosis.
    Returns an in-memory file stream (BytesIO).
    """
    doc = Document()

    # Add logo (if exists)
    logo_path = os.path.join('static', 'images', 'logo.png')
    if os.path.exists(logo_path):
        doc.add_picture(logo_path, width=Inches(1.5))

    doc.add_heading('ProDiag - Prostate Cancer Diagnosis Report', level=1)
    doc.add_paragraph('Empowering Healthcare Through Intelligent Diagnosis\n')

    doc.add_heading('Patient Name', level=2)
    doc.add_paragraph(context.get('name', 'N/A'))

    doc.add_heading('Symptoms Provided', level=2)
    for symptom, value in context.get('symptoms', {}).items():
        doc.add_paragraph(f"{symptom}: {value}")

    doc.add_heading('Diagnosis Result', level=2)
    doc.add_paragraph(context.get('result', 'N/A'))

    doc.add_heading('Confidence Score', level=2)
    doc.add_paragraph(f"{context.get('confidence', 0):.2f}%")

    doc.add_heading('Date of Diagnosis', level=2)
    timestamp = context.get('created_at')
    if timestamp:
        doc.add_paragraph(timestamp.strftime('%B %d, %Y at %I:%M %p'))
    else:
        doc.add_paragraph('N/A')

    doc.add_paragraph()
    doc.add_paragraph(f"© {context.get('current_year', '')} ProDiag. All rights reserved.")

    # Return as in-memory file
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
